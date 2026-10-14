import csv
import io

from openpyxl import load_workbook

from conftest import signup, upload


def _codes(detail: dict) -> set[str]:
    return {r["code"] for r in detail["review_reasons"]}


def test_upload_review_edit_approve_export(client, samples):
    h = signup(client, "ana@empresa-a.test", "Empresa A")

    clean = upload(client, h, samples / "fe_andina_1023.pdf")
    assert clean.status_code == 201
    assert clean.json()["status"] == "approved"  # no issues -> straight through

    bad = upload(client, h, samples / "fe_montana_77.pdf").json()
    assert bad["status"] == "needs_review" and _codes(bad) == {"totals_mismatch"}
    assert {f["name"]: f["source"] for f in bad["fields"]}["total"] == "rules"

    queue = client.get("/api/invoices?status=needs_review", headers=h).json()
    assert [i["id"] for i in queue] == [bad["id"]]

    # Approving with open errors is refused...
    assert client.post(f"/api/invoices/{bad['id']}/approve", headers=h).status_code == 409

    # ...until a human corrects the total (entered in Colombian format).
    fields = {f["name"]: f["value"] for f in bad["fields"]}
    fixed_total = f"{float(fields['subtotal']) + float(fields['iva']):,.2f}"
    fixed_total = fixed_total.replace(",", "X").replace(".", ",").replace("X", ".")
    edited = client.patch(
        f"/api/invoices/{bad['id']}/fields",
        headers=h,
        json={"fields": {"total": f"$ {fixed_total}"}},
    ).json()
    assert _codes(edited) == set()
    total_field = next(f for f in edited["fields"] if f["name"] == "total")
    assert total_field["source"] == "human" and total_field["confidence"] == 1.0

    approved = client.post(f"/api/invoices/{bad['id']}/approve", headers=h).json()
    assert approved["status"] == "approved" and approved["approved_by"] == "ana@empresa-a.test"
    edit_after = client.patch(
        f"/api/invoices/{bad['id']}/fields", headers=h, json={"fields": {"total": "1"}}
    )
    assert edit_after.status_code == 409

    xlsx = client.get("/api/invoices/export?format=xlsx", headers=h)
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("Número", "Fecha", "Proveedor") and len(rows) == 3
    assert {r[0] for r in rows[1:]} == {"FE-1023", "CM-77"}

    text = client.get("/api/invoices/export?format=csv", headers=h).content.decode("utf-8-sig")
    assert len(list(csv.reader(io.StringIO(text)))) == 3


def test_duplicate_flagged_and_owner_can_force_approve(client, samples):
    owner = signup(client, "ana@empresa-a.test", "Empresa A")
    upload(client, owner, samples / "fv_tecno_0458.pdf")
    dup = upload(client, owner, samples / "fv_tecno_0458_dup.pdf").json()
    assert _codes(dup) == {"duplicate"}

    client.post(
        "/api/users",
        headers=owner,
        json={"email": "luis@empresa-a.test", "password": "password123"},
    )
    token = client.post(
        "/api/auth/login", json={"email": "luis@empresa-a.test", "password": "password123"}
    ).json()["access_token"]
    member = {"Authorization": f"Bearer {token}"}
    forced = client.post(f"/api/invoices/{dup['id']}/approve", headers=member, json={"force": True})
    assert forced.status_code == 409  # members cannot override validation
    forced = client.post(f"/api/invoices/{dup['id']}/approve", headers=owner, json={"force": True})
    assert forced.json()["status"] == "approved"


def test_off_template_document_goes_to_review(client, samples):
    h = signup(client, "ana@empresa-a.test")
    detail = upload(client, h, samples / "ds_creativos_981.pdf").json()
    assert detail["status"] == "needs_review" and "missing_field" in _codes(detail)


def test_rejects_non_pdf_and_invalid_edits(client, samples):
    h = signup(client, "ana@empresa-a.test")
    res = client.post("/api/invoices", headers=h, files={"file": ("x.pdf", b"hello", "text/plain")})
    assert res.status_code == 415
    inv = upload(client, h, samples / "fe_caribe_220.pdf").json()
    assert _codes(inv) == {"nit_invalid"}
    original = client.get(f"/api/invoices/{inv['id']}/file", headers=h)
    assert original.headers["content-type"] == "application/pdf"
    assert original.content.startswith(b"%PDF")
    bad = client.patch(
        f"/api/invoices/{inv['id']}/fields", headers=h, json={"fields": {"issue_date": "mañana"}}
    )
    assert bad.status_code == 422
    unknown = client.patch(
        f"/api/invoices/{inv['id']}/fields", headers=h, json={"fields": {"ceo_salary": "1"}}
    )
    assert unknown.status_code == 422


def test_healthz_and_ui_served(client):
    assert client.get("/healthz").json()["llm"] == "fake"
    assert "InvoiceIQ" in client.get("/").text
